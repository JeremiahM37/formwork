import base64
import io
import json
import sys
import time
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'server'))
import _state
from dashboard import documents, model, store
try:
    from fastapi.testclient import TestClient
    from dashboard.app import app
    from docx import Document
    from pypdf import PdfWriter
    HAVE_API=True
except ImportError:
    HAVE_API=False

@unittest.skipUnless(HAVE_API,'install server/requirements.txt')
class DocumentCareerTests(unittest.TestCase):
    def setUp(self):
        store.init()
        with store.db() as db:
            db.execute('DELETE FROM applications');db.execute('DELETE FROM resume_versions')
        self.client=TestClient(app)
        self.record=self.client.post('/api/applications',json={'url':'https://example.test/role','company':'Example','title':'Engineer'}).json()
        self.base=f"/api/applications/{self.record['id']}"

    def test_docx_import_table_order_and_immutable_versions(self):
        doc=Document();doc.add_paragraph('Dana Rivera');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Python';table.cell(0,1).text='5 years';doc.add_paragraph('Education')
        out=io.BytesIO();doc.save(out)
        result=self.client.post('/api/resumes/import',json={'name':'resume.docx','data':base64.b64encode(out.getvalue()).decode()})
        self.assertEqual(result.status_code,200,result.text)
        text=result.json()['text'];self.assertLess(text.index('Python'),text.index('Education'))
        versions=[]
        for content in [text,text+'\nUpdated']:
            versions.append(self.client.post('/api/resumes',json={'name':'Platform','text':content}).json())
        self.assertEqual(len(self.client.get('/api/resumes').json()['versions']),2)
        self.assertEqual(self.client.get('/api/resumes/'+str(versions[0]['id'])).json()['text'],text)
        raw=self.client.get('/api/resumes/'+str(versions[0]['id'])+'/export.docx').content
        self.assertIn('Python', '\n'.join(p.text for p in Document(io.BytesIO(raw)).paragraphs))
        result=self.client.post(self.base+'/resume-version',json={'version_id':versions[0]['id']})
        self.assertEqual(result.status_code,200)
        store.update_application(self.record['id'],status='submitted')
        self.assertEqual(self.client.post(self.base+'/resume-version',json={'version_id':versions[1]['id']}).status_code,409)
        store.update_application(self.record['id'],status='withdrawn')
        self.assertEqual(self.client.post(self.base+'/resume-version',json={'version_id':versions[1]['id']}).status_code,409)

    def test_scanned_and_invalid_input_are_actionable_errors(self):
        pdf=PdfWriter();pdf.add_blank_page(width=600,height=800);out=io.BytesIO();pdf.write(out)
        result=self.client.post('/api/resumes/import',json={'name':'scan.pdf','data':base64.b64encode(out.getvalue()).decode()})
        self.assertEqual(result.status_code,400);self.assertIn('OCR',result.text)
        self.assertEqual(self.client.post('/api/resumes/import',json={'name':'x.docx','data':'bad base64!'}).status_code,400)

    def test_offer_costs_and_one_time_bonus_are_separate(self):
        result=self.client.put(self.base+'/offer',json={'base':100000,'bonus':10000,'equity':20000,'annual_costs':5000,'signing':15000}).json()['offers'][0]
        self.assertEqual(result['recurring'],125000);self.assertEqual(result['first_year'],140000)

    def test_career_drafts_are_saved_and_never_send_or_submit(self):
        from dashboard import career
        with patch.object(career.model,'complete',new=AsyncMock(return_value='I am writing to apply. I built Python tools.')) as complete:
            result=self.client.post(self.base+'/drafts',json={'kind':'followup'}).json()
            self.assertTrue(result['needsApproval']);self.assertTrue(result['tells'])
            self.assertIn('TIMELINE',complete.call_args.args[0][1]['content'])
        self.assertEqual(len(self.client.get(self.base+'/drafts').json()['drafts']),1)
        self.assertEqual(store.get_application(self.record['id'])['status'],'queued')
        self.assertEqual(self.client.post(self.base+'/drafts',json={'kind':'feedback'}).status_code,400)

    def test_provider_keys_are_private_and_never_reused_for_new_destination(self):
        with patch.object(model,'PROVIDER_PATH',Path(_state.STATE_DIR)/'test-provider.json'):
            model.save_config({'provider':'openai-compatible','url':'https://first.test/v1','model':'test','key':'fixture-secret'})
            self.assertNotIn('fixture-secret',json.dumps(model.public_config()))
            self.assertEqual(model.PROVIDER_PATH.stat().st_mode & 0o777,0o600)
            model.save_config({'provider':'openai-compatible','url':'https://second.test/v1','model':'test','key':''})
            self.assertFalse(model.public_config()['hasKey'])

    def test_saved_answers_merge_into_profile_and_browser_sync_receipt(self):
        from dashboard import app as dashboard
        path=Path(_state.STATE_DIR)/'profile-sync.json';path.write_text(json.dumps({'identity':{'full_name':'Dana'},'work_authorization':{'requires_sponsorship_now':True}}))
        with patch.object(dashboard,'PROFILE_JSON',path),patch.object(dashboard.browser,'run',new=AsyncMock(return_value={'ok':True})) as sync:
            result=self.client.post('/api/profile/answers',json={'text':json.dumps({'work_authorization':{'requires_sponsorship_now':False}})}).json()
            self.assertTrue(result['synced'])
            self.assertFalse(sync.call_args.args[1]['profile']['work_authorization']['requires_sponsorship_now'])
            self.assertEqual(json.loads(path.read_text())['identity']['full_name'],'Dana')

    def test_provider_transports_over_real_local_http(self):
        import asyncio
        import threading
        from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
        received=[]
        class Handler(BaseHTTPRequestHandler):
            def do_POST(self):
                body=json.loads(self.rfile.read(int(self.headers['Content-Length'])))
                received.append((self.path,body,dict(self.headers)))
                if self.path.endswith('/api/chat'): response={'message':{'content':'ollama ready'}}
                elif self.path.endswith('/messages'): response={'content':[{'type':'text','text':'anthropic ready'}]}
                elif self.path.endswith('/chat/completions'): response={'choices':[{'message':{'content':'compatible ready'}}]}
                else: response={'content':'homelab ready'}
                payload=json.dumps(response).encode();self.send_response(200);self.send_header('Content-Length',str(len(payload)));self.end_headers();self.wfile.write(payload)
            def log_message(self,*args): pass
        server=ThreadingHTTPServer(('127.0.0.1',0),Handler);worker=threading.Thread(target=server.serve_forever,daemon=True);worker.start()
        try:
            with patch.object(model,'PROVIDER_PATH',Path(_state.STATE_DIR)/'wire-provider.json'):
                for provider in ['homelab','ollama','openai-compatible','anthropic']:
                    model.save_config({'provider':provider,'url':f'http://127.0.0.1:{server.server_port}','model':'fixture','key':'test-key'})
                    result=asyncio.run(model.complete([{'role':'system','content':'Follow source facts.'},{'role':'user','content':'Ready?'}],want_json=True))
                    self.assertIn('ready',result)
                self.assertEqual(received[1][1]['think'],False)
                self.assertIn('system',received[3][1])
        finally:
            server.shutdown();server.server_close();worker.join()

    def test_requirement_analysis_refuses_invented_quotes(self):
        from dashboard.career import validate_fit_analysis
        result=validate_fit_analysis({'requirements':[
            {'requirement':'Python and SQL experience','evidence':'Built Python ingestion services','importance':'required','assessment':'met'},
            {'requirement':'PhD in Robotics','evidence':'Earned a PhD','importance':'required','assessment':'met'},
            {'requirement':'Kubernetes experience','evidence':'Managed Kubernetes for ten years','assessment':'met'}]},
            'Python and SQL experience. Kubernetes experience.', 'Built Python ingestion services. I use SQL.')
        self.assertEqual(len(result['requirements']),2)
        self.assertEqual(result['requirements'][1]['assessment'],'unverified')
        self.assertEqual(result['requirements'][1]['evidence'],'')
