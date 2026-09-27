import io
import json
import sys
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'server'))
import _state
from dashboard import documents, latex, resume_layout as layout
from dashboard.app import app
from fastapi.testclient import TestClient
from docx import Document
from pypdf import PdfReader


class ResumeLayoutTests(unittest.TestCase):
    def setUp(self):
        self.document=layout.Layout(name='Dana Rivéra',contact='dana@example.test',template='compact',paper='a4',font_size=10,margin=.6,
            sections=[layout.Section(heading='Skills',text='Python, SQL'),
                layout.Section(heading='Experience',entries=[layout.Entry(title='Software Engineer',subtitle='Fixture Company',dates='2022 – Present',
                    bullets=['Reduced latency by 40%.','Built A&B tools with $5 costs.'])])])
        self.client=TestClient(app)

    def test_word_layout_preserves_order_measurements_and_page_settings(self):
        doc=Document(io.BytesIO(layout.word(self.document)))
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertLess(text.index('Skills'),text.index('Experience'))
        self.assertIn('Reduced latency by 40%.',text)
        self.assertIn('A&B tools with $5',text)
        self.assertAlmostEqual(doc.sections[0].left_margin.inches,.6,places=3)
        self.assertAlmostEqual(doc.sections[0].page_width.inches,8.2677,places=3)
        self.assertEqual(doc.styles['Normal'].font.size.pt,10)

    def test_real_pdf_is_readable_and_entries_remain_tailorable(self):
        saved=self.client.post('/api/resumes',json={'name':'Layout test','format':'structured','text':self.document.model_dump_json()})
        self.assertEqual(saved.status_code,200,saved.text)
        version=saved.json()
        index=latex.parse(documents.version_tex(version))
        self.assertEqual(index.entries[0].title,'Software Engineer')
        self.assertEqual(len(index.entries[0].bullets),2)
        pdf=self.client.get(f"/api/resumes/{version['id']}/export.pdf?inline=1")
        self.assertEqual(pdf.status_code,200,pdf.text if pdf.status_code!=200 else '')
        self.assertTrue(pdf.headers['content-disposition'].startswith('inline'))
        reader=PdfReader(io.BytesIO(pdf.content))
        text='\n'.join(page.extract_text() for page in reader.pages)
        self.assertEqual(len(reader.pages),1)
        self.assertIn('40%',text)
        self.assertIn('Rivéra',text)
        self.assertIn('2022 – Present',text)
        self.assertLess(text.index('Skills'),text.index('Experience'))
        self.assertIn('Fixture Company',text)

    def test_invalid_layout_is_rejected_before_persistence(self):
        invalid=self.document.model_dump();invalid['margin']=0
        result=self.client.post('/api/resumes',json={'name':'Invalid','format':'structured','text':json.dumps(invalid)})
        self.assertEqual(result.status_code,400)

    def test_profile_seed_preserves_gpa_honors_dates_and_project_links(self):
        profile={'identity':{'full_name':'Dana'},'education':[{'school':'Blue Ridge','degree':'B.S.','field_of_study':'Computer Science','gpa':'3.2','honors':['Dean list'],'start':'2018','end':'2022'}],
            'projects':[{'name':'Tool','url':'https://example.test/tool'}],'skills':{'languages':['Python']}}
        result=layout.from_profile(profile)
        self.assertIn('GPA: 3.2',result.sections[0].entries[0].bullets)
        self.assertIn('Dean list',result.sections[0].entries[0].bullets)
        self.assertEqual(result.sections[0].entries[0].dates,'2018 – 2022')
        self.assertIn('https://example.test/tool',result.sections[1].entries[0].bullets)

    def test_all_eight_layouts_export_readable_pdf_and_word_with_same_facts(self):
        for template in layout.TEMPLATES:
            with self.subTest(template=template):
                doc=self.document.model_copy(update={'template':template})
                tex=layout.tex(doc)
                self.assertEqual(latex.parse(tex).entries[0].title,'Software Engineer')
                pdf=documents.resume.compile_pdf(tex,'template-review-'+template)
                text='\n'.join(p.extract_text() for p in PdfReader(str(pdf)).pages)
                word=Document(io.BytesIO(layout.word(doc)))
                word_text='\n'.join(p.text for p in word.paragraphs)
                for fact in ['Rivéra','40%','Fixture Company','2022 – Present']:
                    self.assertIn(fact,text);self.assertIn(fact,word_text)
                self.assertLess(text.index('Skills'),text.index('Experience'))
