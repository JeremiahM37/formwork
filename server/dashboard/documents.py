"""Resume import, immutable versions and editable Word exports."""
import base64
import binascii
import io
import re
import time
import zipfile
from typing import Literal

from fastapi import APIRouter, HTTPException
from fastapi.responses import Response, FileResponse
from pydantic import BaseModel, Field

from . import latex, resume, resume_layout, store, tracking, matching
from pydantic import ValidationError

router = APIRouter(prefix="/api")
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ImportIn(BaseModel):
    name: str = Field(min_length=1, max_length=300)
    data: str = Field(max_length=14000000)


class VersionIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    text: str = Field(min_length=1, max_length=100000)
    format: Literal["text", "tex", "structured"] = "text"


class SelectionIn(BaseModel):
    version_id: int


def extract(name, raw):
    suffix = name.lower().rsplit('.', 1)[-1]
    if len(raw) > 10 * 1024 * 1024:
        raise ValueError("Document exceeds 10 MB")
    if suffix == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            raise ValueError("Unlock the PDF before importing it")
        if len(reader.pages) > 30:
            raise ValueError("Import a resume of at most 30 pages")
        text = "\n\n".join(p.extract_text() or "" for p in reader.pages)
    elif suffix == "docx":
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Expanded document exceeds 30 MB")
        doc = Document(io.BytesIO(raw))
        # Preserve paragraph/table order; table-heavy resumes are common.
        parts = []
        for block in doc.iter_inner_content():
            if hasattr(block, "text"):
                parts.append(block.text)
            else:
                parts.extend(" | ".join(c.text for c in row.cells) for row in block.rows)
        text = "\n".join(parts)
    elif suffix in {"txt", "md", "tex"}:
        text = raw.decode("utf-8-sig")
    else:
        raise ValueError("Use PDF, DOCX, TXT, Markdown or LaTeX")
    if not text.strip():
        raise ValueError("No text found. For a scanned PDF, use OCR or paste the resume text.")
    if len(text) > 100000:
        raise ValueError("Extracted resume exceeds 100,000 characters")
    return text


@router.post("/resumes/import")
def import_resume(body: ImportIn):
    try:
        text = extract(body.name, base64.b64decode(body.data, validate=True))
    except (ValueError, binascii.Error, zipfile.BadZipFile) as err:
        raise HTTPException(400, str(err))
    except Exception as err:
        raise HTTPException(400, f"Could not read this document: {type(err).__name__}")
    return {"text":text, "format":"tex" if body.name.lower().endswith('.tex') else "text",
            "note":"Review the extracted text before saving. Import does not change your profile or existing resumes."}


@router.get("/resumes")
def versions():
    return {"versions":tracking.rows("SELECT id,name,format,created_at FROM resume_versions ORDER BY id DESC")}


def get_version(version_id):
    rows = tracking.rows("SELECT * FROM resume_versions WHERE id=?", (version_id,))
    if not rows:
        raise HTTPException(404, "no such resume version")
    return rows[0]


@router.get("/resumes/{version_id}")
def version(version_id: int):
    return get_version(version_id)


@router.post("/resumes")
def save_version(body: VersionIn):
    if body.format == "structured":
        try:
            body.text = resume_layout.Layout.model_validate_json(body.text).model_dump_json()
        except ValidationError:
            raise HTTPException(400, "Invalid resume layout; check the sections and layout settings")
    with store.db() as conn:
        cur = conn.execute("INSERT INTO resume_versions(name,text,format,created_at) VALUES(?,?,?,?)",
                           (body.name,body.text,body.format,time.time()))
        version_id = cur.lastrowid
    return get_version(version_id)


@router.post("/applications/{app_id}/resume-version")
def select_version(app_id: int, body: SelectionIn):
    # Reuse the same preparation/submission boundary as field and letter edits.
    from .app import _editable_application
    record = _editable_application(app_id)
    selected = get_version(body.version_id)
    snapshot = {**record["snapshot"], "resumeVersion":selected["id"]}
    store.update_application(app_id, snapshot=snapshot)
    return {"selected":selected["name"], "note":"Prepare the application to generate and attach this version."}


def version_tex(item):
    if item["format"] == "structured":
        return resume_layout.tex(resume_layout.Layout.model_validate_json(item["text"]))
    if item["format"] == "tex":
        return item["text"]
    paragraphs = "\n\n".join("\\noindent " + resume._tex_escape(line) for line in item["text"].splitlines() if line.strip())
    return r"\documentclass[10pt]{article}\usepackage[margin=0.7in]{geometry}\usepackage[T1]{fontenc}\pagestyle{empty}\begin{document}" + "\n" + paragraphs + "\n" + r"\end{document}"


def word_bytes(text):
    from docx import Document
    from docx.shared import Inches, Pt
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(.7)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    for line in text.splitlines():
        if not line.strip():
            continue
        if line.lstrip().startswith(('- ', '• ')):
            doc.add_paragraph(line.lstrip()[2:], style='List Bullet')
        else:
            doc.add_paragraph(line)
    out = io.BytesIO(); doc.save(out)
    return out.getvalue()


@router.get("/resumes/{version_id}/export.docx")
def export_version(version_id: int):
    item = get_version(version_id)
    if item["format"] == "structured":
        return Response(resume_layout.word(resume_layout.Layout.model_validate_json(item["text"])), media_type=DOCX_MIME,
                        headers={"Content-Disposition": f'attachment; filename="resume-{version_id}.docx"'})
    text = latex.document_text(item["text"]) if item["format"] == "tex" else item["text"]
    return Response(word_bytes(text), media_type=DOCX_MIME, headers={"Content-Disposition": f'attachment; filename="resume-{version_id}.docx"'})


@router.get("/resumes/{version_id}/export.pdf")
def export_pdf(version_id: int, inline: bool = False):
    try:
        path = resume.compile_pdf(version_tex(get_version(version_id)), f"resume-{version_id}")
    except RuntimeError as err:
        raise HTTPException(400, str(err))
    return FileResponse(path, media_type="application/pdf", filename=f"resume-{version_id}.pdf", content_disposition_type="inline" if inline else "attachment")


@router.get("/resume-builder/profile")
def builder_profile():
    from .app import _profile
    return resume_layout.from_profile(_profile()).model_dump()


@router.get("/applications/{app_id}/cover.docx")
def export_cover(app_id: int):
    record = tracking.application(app_id)
    text = record["snapshot"].get("cover", {}).get("text", "")
    if not text:
        raise HTTPException(404, "no cover letter drafted")
    return Response(word_bytes(text), media_type=DOCX_MIME, headers={"Content-Disposition": f'attachment; filename="cover-{app_id}.docx"'})


@router.get("/applications/{app_id}/resume.docx")
def export_application_resume(app_id: int):
    record = tracking.application(app_id)
    text = record["snapshot"].get("resumeText", "")
    if not text:
        raise HTTPException(404, "prepare the application to save its resume snapshot")
    return Response(word_bytes(text), media_type=DOCX_MIME, headers={"Content-Disposition": f'attachment; filename="resume-application-{app_id}.docx"'})


class ResumeComparisonIn(BaseModel):
    version_ids: list[int] = Field(min_length=2, max_length=10)


def version_text(item):
    if item["format"] == "structured":
        doc = resume_layout.Layout.model_validate_json(item["text"])
        return "\n".join([doc.headline, *[
            "\n".join([section.heading, section.text, *[
                "\n".join([entry.title, entry.subtitle, *entry.bullets])
                for entry in section.entries]]) for section in doc.sections]])
    return latex.document_text(item["text"]) if item["format"] == "tex" else item["text"]


@router.post("/applications/{app_id}/compare-resumes")
def compare_resumes(app_id: int, body: ResumeComparisonIn):
    record = tracking.application(app_id)
    if len(set(body.version_ids)) != len(body.version_ids):
        raise HTTPException(400, "Choose distinct resume versions")
    results = []
    for version_id in body.version_ids:
        item = get_version(version_id)
        # Compare only what each document says. Shared profile skills would
        # otherwise make a weak resume look as complete as a targeted version.
        fit = matching.fit(record["snapshot"].get("posting", ""), {}, version_text(item))
        results.append({"id": item["id"], "name": item["name"], "fit": fit})
    results.sort(key=lambda r: -(r["fit"]["coverage"] if r["fit"]["coverage"] is not None else -1))
    return {"versions": results, "selected": record["snapshot"].get("resumeVersion"),
            "note": "Compares recognized skill evidence in each document only. This is not an ATS score or a hiring prediction. Nothing is selected or edited automatically."}


class PostingAnalysisIn(BaseModel):
    description: str = Field(min_length=1, max_length=30000)
    version_id: int | None = None


def sponsorship_excerpts(description):
    """Quote posting language without inferring eligibility or employer history."""
    sentences = re.split(r'(?<=[.!?])\s+|\n+', description)
    pattern = re.compile(r'\b(?:visa|h[- ]?1b|work authori[sz]ation|authori[sz]ed to work|immigration|sponsorship)\b', re.I)
    return list(dict.fromkeys(s.strip() for s in sentences if pattern.search(s)))[:8]


@router.post("/posting-analysis")
def analyze_posting(body: PostingAnalysisIn):
    from .app import _profile, _read, RESUME_TEX
    if body.version_id is not None:
        selected = get_version(body.version_id)
        fit = matching.fit(body.description, {}, version_text(selected))
        label = selected['name']
    else:
        fit = matching.fit(body.description, _profile(), latex.document_text(_read(RESUME_TEX)))
        label = 'Saved profile and master resume'
    return {"fit": fit, "resume": label, "versions": versions()['versions'],
            "sponsorshipExcerpts": sponsorship_excerpts(body.description),
            "note": "Analysis only: no form fields were changed and no application was saved."}
